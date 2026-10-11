import io
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from motor import excel
from motor.banco import Ejercicio


def ej(comprobaciones, hojas=None, **kw) -> Ejercicio:
    hojas = hojas or [{
        "nombre": "Sueldos",
        "columnas": ["Empleado", "USD", "EUR"],
        "filas": [["Ana", 1000, None], ["Luis", 2000, None], ["Eva", 1500, None]],
        "celdas": {"E1": "Tipo de cambio", "F1": 0.5},
    }]
    base = dict(id="xl-t", tipo="excel", tema="excel_pq", concepto="c", nivel="basico", titulo="t", enunciado="e",
                origen="o", nombre_archivo="t.xlsx", hojas=hojas, comprobaciones=comprobaciones)
    base.update(kw)
    return Ejercicio(**base)


def guardar(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------- comparador
@pytest.mark.parametrize("obtenido, esperado, razon", [
    (10.004, 10, None),
    (10.5, 10, "distinto"),
    ("10", 10, "texto"),
    (None, 10, "vacio"),
    (True, 1, "distinto"),
    ("Alto", "Alto", None),
    ("alto", "Alto", "distinto"),
    ("Alto ", "Alto", "distinto"),
    (3, "3", "numero"),
    (None, None, None),
    ("", None, None),
    (5, None, "distinto"),
    (datetime(2014, 3, 1, 10, 30), date(2014, 3, 1), None),
    (datetime(2014, 3, 2), date(2014, 3, 1), "distinto"),
    (41699, date(2014, 3, 1), None),          # numero de serie de Excel
    ("01/03/2014", date(2014, 3, 1), "texto"),
])
def test_razon_distinta(obtenido, esperado, razon):
    assert excel.razon_distinta(obtenido, esperado) == razon


def test_ignorar_mayusculas():
    assert excel.razon_distinta("ALTO", "Alto", ignorar_mayusculas=True) is None


# -------------------------------------------------------------- libro de partida
def test_el_libro_de_partida_tiene_los_datos_y_las_celdas_extra():
    e = ej([])
    wb = load_workbook(io.BytesIO(excel.generar_libro(e)))
    ws = wb["Sueldos"]
    assert [c.value for c in ws[1]][:3] == ["Empleado", "USD", "EUR"]
    assert ws["B3"].value == 2000
    assert ws["F1"].value == 0.5


def test_csv_con_consulta_columnas_y_tipos(tmp_path):
    (tmp_path / "d.csv").write_text("a;b;c\n1;x; 2.50 \n2;y; 3.00 \n3;x; 4.00 \n", encoding="utf-8")
    e = ej([], hojas=[{"nombre": "H", "csv": "d.csv", "sep": ";", "consulta": "b == 'x'", "elegir": ["a", "c"],
                       "numericas": ["a"]}])
    libro = load_workbook(io.BytesIO(excel.generar_libro(e, tmp_path)))["H"]
    assert [[c.value for c in f] for f in libro.iter_rows()] == [["a", "c"], [1, " 2.50 "], [3, " 4.00 "]]


def test_csv_que_falta_da_error_claro(tmp_path):
    e = ej([], hojas=[{"nombre": "H", "csv": "no_existe.csv"}])
    r = excel.corregir_excel(e, excel.generar_libro(ej([])), tmp_path)
    assert r["estado"] == "error_banco" and "preparar_datos" in r["mensajes"][0]


def test_calc_solo_permite_expresiones_sobre_los_datos():
    e = ej([])
    datos = excel.datos_origen(e)
    assert excel.evaluar("datos['Sueldos']['USD'].sum()", datos) == 4500
    assert excel.evaluar("(datos['Sueldos']['USD'] * 2).tolist()", datos) == [2000, 4000, 3000]
    with pytest.raises(excel.ErrorExcel):
        excel.evaluar("__import__('os').getcwd()", datos)


# -------------------------------------------------------------- celdas y formulas
COMPROB_ABS = [
    {"tipo": "columna", "hoja": "Sueldos", "encabezado": "EUR", "calc": "(datos['Sueldos']['USD'] * 0.5).tolist()",
     "formula": True, "formula_contiene": ["$F$1"]},
]


def test_el_libro_sin_resolver_no_pasa_y_el_resuelto_si():
    e = ej(COMPROB_ABS)
    assert excel.corregir_excel(e, excel.generar_libro(e))["estado"] == "fallo"
    r = excel.corregir_excel(e, excel.libro_resuelto(e))
    assert r["estado"] == "ok", r["mensajes"]


def test_formula_sin_resultado_guardado_explica_que_hay_que_guardar_en_excel():
    e = ej(COMPROB_ABS)
    sin = excel.sin_calcular(excel.libro_resuelto(e))
    # sin_calcular reescribe con openpyxl: las formulas siguen, los valores guardados no
    r = excel.corregir_excel(e, sin)
    assert r["estado"] == "fallo"
    assert "Ctrl+G" in r["mensajes"][0]


def test_valor_escrito_a_mano_donde_se_pide_formula():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sueldos"
    for j, h in enumerate(["Empleado", "USD", "EUR"], start=1):
        ws.cell(row=1, column=j, value=h)
    for i, (n, u) in enumerate([("Ana", 1000), ("Luis", 2000), ("Eva", 1500)], start=2):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=u)
        ws.cell(row=i, column=3, value=u * 0.5)
    r = excel.corregir_excel(ej(COMPROB_ABS), guardar(wb))
    assert r["estado"] == "fallo" and "valores escritos" in r["mensajes"][0]


def test_la_formula_debe_contener_la_referencia_absoluta():
    e = ej([{**COMPROB_ABS[0], "formula_contiene": ["$Z$9"]}])
    r = excel.corregir_excel(e, excel.libro_resuelto(ej(COMPROB_ABS)))
    assert r["estado"] == "fallo" and "$Z$9" in r["mensajes"][0]


def test_celda_con_valor_y_tolerancia_y_texto_en_vez_de_numero():
    e = ej([{"tipo": "celda", "hoja": "Sueldos", "celda": "H1", "calc": "datos['Sueldos']['USD'].sum()"}])
    for valor, ok in [(4500, True), (4500.004, True), (4400, False), ("4500", False), (None, False)]:
        wb = load_workbook(io.BytesIO(excel.generar_libro(e)))
        wb["Sueldos"]["H1"] = valor
        r = excel.corregir_excel(e, guardar(wb))
        assert (r["estado"] == "ok") is ok, (valor, r["mensajes"])
    wb = load_workbook(io.BytesIO(excel.generar_libro(e)))
    wb["Sueldos"]["H1"] = "4500"
    assert "texto" in excel.corregir_excel(e, guardar(wb))["mensajes"][0]


def test_los_mensajes_no_revelan_el_valor_esperado():
    e = ej([{"tipo": "celda", "hoja": "Sueldos", "celda": "H1", "valor": 987654.321}])
    wb = load_workbook(io.BytesIO(excel.generar_libro(e)))
    wb["Sueldos"]["H1"] = 1
    assert "987654" not in " ".join(excel.corregir_excel(e, guardar(wb))["mensajes"])


def test_hoja_con_otro_nombre_en_mayusculas_se_acepta():
    e = ej([{"tipo": "celda", "hoja": "Resumen", "celda": "A1", "valor": 3}])
    wb = Workbook()
    wb.active.title = "resumen "
    wb.active["A1"] = 3
    assert excel.corregir_excel(e, guardar(wb))["estado"] == "ok"


def test_fechas_como_datetime_y_como_serie():
    e = ej([{"tipo": "celda", "hoja": "H", "celda": "A1", "valor": date(2014, 3, 1)}])
    for v in (datetime(2014, 3, 1), 41699):
        wb = Workbook()
        wb.active.title = "H"
        wb.active["A1"] = v
        assert excel.corregir_excel(e, guardar(wb))["estado"] == "ok"


# -------------------------------------------------------------- columnas
def _libro_columna(valores, enc="Total") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = enc
    for i, v in enumerate(valores, start=2):
        ws.cell(row=i, column=1, value=v)
    return guardar(wb)


def test_columna_en_orden_y_fila_del_primer_error():
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Total", "valores": [1, 2, 3]}])
    assert excel.corregir_excel(e, _libro_columna([1, 2, 3]))["estado"] == "ok"
    r = excel.corregir_excel(e, _libro_columna([1, 9, 3]))
    assert r["estado"] == "fallo" and "fila 3" in r["mensajes"][0]
    r = excel.corregir_excel(e, _libro_columna([1, 2]))
    assert "2 valores" in r["mensajes"][0] and "3" in r["mensajes"][0]


def test_columna_sin_importar_el_orden():
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Total", "valores": [1, 2, 3], "orden": False}])
    assert excel.corregir_excel(e, _libro_columna([3, 1, 2]))["estado"] == "ok"
    assert excel.corregir_excel(e, _libro_columna([3, 1, 1]))["estado"] == "fallo"


def test_encabezado_que_no_existe():
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Otro", "valores": [1]}])
    assert "Otro" in excel.corregir_excel(e, _libro_columna([1]))["mensajes"][0]


def test_columna_solo_numeros_sin_espacios_y_sin_duplicados():
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Total", "solo_numeros": True}])
    assert excel.corregir_excel(e, _libro_columna([1, 2.5]))["estado"] == "ok"
    r = excel.corregir_excel(e, _libro_columna([1, " 2.50 "]))
    assert r["estado"] == "fallo" and "fila 3" in r["mensajes"][0]
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Total", "sin_espacios_extra": True}])
    assert excel.corregir_excel(e, _libro_columna(["ana perez", "luis"]))["estado"] == "ok"
    assert excel.corregir_excel(e, _libro_columna(["ana perez", "luis "]))["estado"] == "fallo"
    assert excel.corregir_excel(e, _libro_columna(["ana  perez"]))["estado"] == "fallo"
    e = ej([{"tipo": "columna", "hoja": "H", "encabezado": "Total", "sin_duplicados": True}])
    assert excel.corregir_excel(e, _libro_columna(["a", "b"]))["estado"] == "ok"
    assert excel.corregir_excel(e, _libro_columna(["a", "b", "a"]))["estado"] == "fallo"


# -------------------------------------------------------------- estructura
def test_encabezados_hojas_y_filas():
    e = ej([
        {"tipo": "encabezados", "hoja": "H", "valores": ["Total"]},
        {"tipo": "hoja", "hoja": "H"},
        {"tipo": "hoja", "hoja": "Basura", "existe": False},
        {"tipo": "filas", "hoja": "H", "n": 2},
    ])
    assert excel.corregir_excel(e, _libro_columna([1, 2]))["estado"] == "ok"
    r = excel.corregir_excel(e, _libro_columna([1, 2, 3]))
    assert r["estado"] == "fallo" and "3 filas" in r["mensajes"][0]


def test_tabla_de_excel_y_celdas_combinadas():
    e = ej([{"tipo": "tabla", "hoja": "H"}, {"tipo": "sin_combinadas", "hoja": "H"}])
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"], ws["A2"] = "x", 1
    ws.merge_cells("C1:D1")
    r = excel.corregir_excel(e, guardar(wb))
    assert [c["ok"] for c in r["comprobaciones"]] == [False, False]
    assert "Insertar > Tabla" in r["mensajes"][0]
    r = excel.corregir_excel(e, excel.libro_resuelto(ej(e.comprobaciones, hojas=[{"nombre": "H", "columnas": ["x"], "filas": [[1]]}])))
    assert r["estado"] == "ok", r["mensajes"]


def test_formato_condicional_por_tipo_operador_y_rango():
    e = ej([{"tipo": "formato_condicional", "hoja": "H", "rango": "A2:A9", "regla": "cellIs", "operador": "greaterThan"}])
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    r = excel.corregir_excel(e, guardar(wb))
    assert r["estado"] == "fallo" and "ninguna regla" in r["mensajes"][0]
    ok = excel.libro_resuelto(ej(e.comprobaciones, hojas=[{"nombre": "H", "columnas": ["x"], "filas": [[1]]}]))
    assert excel.corregir_excel(e, ok)["estado"] == "ok"
    # regla aplicada a otro rango: no vale
    e2 = ej([{**e.comprobaciones[0], "rango": "B2:B9"}])
    r = excel.corregir_excel(e2, ok)
    assert r["estado"] == "fallo" and "cellIs" in r["mensajes"][0]


def test_estructuras_que_openpyxl_no_crea_se_detectan_en_el_zip():
    comps = [
        {"tipo": "grafico", "tipo_grafico": "circular", "etiquetas": "porcentaje"},
        {"tipo": "tabla_dinamica"},
        {"tipo": "power_query", "contiene_m": ["en-US"]},
    ]
    e = ej(comps)
    base = excel.generar_libro(e)
    r = excel.corregir_excel(e, base)
    assert r["estado"] == "fallo" and len(r["mensajes"]) == 3
    resuelto = excel.libro_resuelto(e)
    r = excel.corregir_excel(e, resuelto)
    assert r["estado"] == "ok", r["mensajes"]
    assert "en-US" in excel.codigo_m(excel.Libro(resuelto))


def test_grafico_de_otro_tipo_o_sin_etiquetas_no_vale():
    e = ej([{"tipo": "grafico", "tipo_grafico": "circular", "etiquetas": "porcentaje"}])
    barras = excel.libro_resuelto(ej([{"tipo": "grafico", "tipo_grafico": "barras"}]))
    assert "circular" in excel.corregir_excel(e, barras)["mensajes"][0]
    sin_et = excel.libro_resuelto(ej([{"tipo": "grafico", "tipo_grafico": "circular"}]))
    assert "porcentaje" in excel.corregir_excel(e, sin_et)["mensajes"][0]


def test_las_comprobaciones_blandas_avisan_pero_no_suspenden():
    e = ej([{"tipo": "power_query", "blando": True}, {"tipo": "hoja", "hoja": "Sueldos"}])
    r = excel.corregir_excel(e, excel.generar_libro(e))
    assert r["estado"] == "ok"
    assert r["comprobaciones"][0]["ok"] is False and r["comprobaciones"][0]["blando"] is True


# -------------------------------------------------------------- archivos malos
@pytest.mark.parametrize("contenido", [b"", b"esto no es un excel", b"PK\x03\x04roto"])
def test_archivo_que_no_es_un_xlsx(contenido):
    r = excel.corregir_excel(ej([{"tipo": "hoja", "hoja": "H"}]), contenido)
    assert r["estado"] == "archivo_invalido" and r["mensajes"]


def test_zip_que_no_es_un_libro(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "hola")
    r = excel.corregir_excel(ej([{"tipo": "hoja", "hoja": "H"}]), buf.getvalue())
    assert r["estado"] == "archivo_invalido"


def test_archivo_demasiado_grande(monkeypatch):
    monkeypatch.setattr(excel, "MAX_BYTES", 10)
    r = excel.corregir_excel(ej([{"tipo": "hoja", "hoja": "H"}]), b"x" * 20)
    assert r["estado"] == "archivo_invalido" and "15 MB" in r["mensajes"][0]


def test_columna_vacia_se_dice_con_claridad_y_no_como_valores_a_mano():
    e = ej(COMPROB_ABS)
    r = excel.corregir_excel(e, excel.generar_libro(e))
    assert r["estado"] == "fallo"
    assert "está vacía" in r["mensajes"][0]
    assert "valores escritos" not in r["mensajes"][0]


def test_archivo_roto_da_un_mensaje_amable():
    import io as _io
    import zipfile as _zf
    buf = _io.BytesIO()
    with _zf.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<roto")
        z.writestr("xl/workbook.xml", "<roto")
    r = excel.corregir_excel(ej([{"tipo": "hoja", "hoja": "H"}]), buf.getvalue())
    assert r["estado"] == "archivo_invalido" and "No he podido abrir el archivo" in r["mensajes"][0]
